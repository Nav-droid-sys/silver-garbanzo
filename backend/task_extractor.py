# task_extractor.py - Extract Task IDs & Billing Details to Master Excel File
import os
import sys
from datetime import datetime, timedelta
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    try: sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except: pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Resolve paths dynamically relative to workspace root
BACKEND_DIR = Path(__file__).resolve().parent
BASE_DIR = BACKEND_DIR.parent

DEFAULT_INPUT_NAME = "abhinav-input-Jun'26dump.xlsx"
SOURCE = BASE_DIR / DEFAULT_INPUT_NAME
MASTER_FILE_NAME = "Billing_Extracted_Report.xlsx"
MASTER_PATH = BASE_DIR / MASTER_FILE_NAME

# Sheet layout configuration (exclusively 'Jira Dump')
SHEETS = {
    "Jira Dump": {
        "hdr": 7, "start": 8, "task": 6, "name": 7, "points": 12,
        "milestone": 10, "status": 3, "dac_month": 9, "deliv_date": 8
    },
}


def cell_val(cell):
    """Return cell value, skip formulas."""
    v = cell.value
    return None if isinstance(v, str) and v.startswith("=") else v


def extract_task_date(date_val):
    """Return datetime.datetime object or None from date cell value."""
    if date_val is None:
        return None
    if isinstance(date_val, datetime):
        return date_val
    if hasattr(date_val, "year") and hasattr(date_val, "month") and hasattr(date_val, "day"):
        return datetime(date_val.year, date_val.month, date_val.day)
    if isinstance(date_val, str):
        v = date_val.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m"):
            try:
                return datetime.strptime(v, fmt)
            except ValueError:
                pass
    return None


def format_month_year(date_val):
    """Convert datetime/date or string into 'Mon YYYY' format (e.g. 'Jan 2025')."""
    dt = extract_task_date(date_val)
    if dt:
        return dt.strftime("%b %Y")
    return "Unknown Date"


def parse_month_year_sortkey(my_str):
    """Sort key for 'Mon YYYY' strings."""
    try:
        return datetime.strptime(my_str, "%b %Y")
    except ValueError:
        return datetime.min


def parse_user_month_year(s):
    """Parse user entered month & year string into datetime object (first day of month)."""
    if not s:
        return None
    s = s.strip().title()
    for fmt in ("%b %Y", "%m/%Y", "%Y-%m", "%b-%Y", "%B %Y", "%b %y", "%B %y", "%m-%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def load_data(source_path=None):
    """Load workbook and collect all assignees with their task counts from Jira Dump."""
    target_src = Path(source_path) if source_path else SOURCE
    if not target_src.exists():
        # Fallback: search for any .xlsx file in BASE_DIR
        xlsx_files = list(BASE_DIR.glob("*.xlsx"))
        xlsx_files = [f for f in xlsx_files if not f.name.startswith("Billing_Extracted_Report")]
        if xlsx_files:
            target_src = xlsx_files[0]
        else:
            print(f"  ERROR: Source file {target_src} not found"); sys.exit(1)

    print(f"  Loading {target_src.name} ...")
    wb = openpyxl.load_workbook(str(target_src), data_only=True)
    names = {}  # name -> count

    for sname, cfg in SHEETS.items():
        if sname not in wb.sheetnames: continue
        ws = wb[sname]
        for r in range(cfg["start"], ws.max_row + 1):
            v = cell_val(ws.cell(r, cfg["name"]))
            if v and isinstance(v, str):
                n = v.strip()
                if n and n != "-":
                    names[n] = names.get(n, 0) + 1

    return wb, names, target_src


def extract_task_details(wb, target):
    """Extract Task IDs, Estimator Points, Milestone, Status & Month/Year for target name from Jira Dump."""
    tasks = []
    for sname, cfg in SHEETS.items():
        if sname not in wb.sheetnames: continue
        ws = wb[sname]
        for r in range(cfg["start"], ws.max_row + 1):
            v = cell_val(ws.cell(r, cfg["name"]))
            if not (v and isinstance(v, str) and v.strip().lower() == target.lower()):
                continue
            
            tid = cell_val(ws.cell(r, cfg["task"]))
            if tid is None: continue
            
            pts = cell_val(ws.cell(r, cfg["points"]))
            try:
                pts_val = float(pts) if pts is not None else 0.0
            except (ValueError, TypeError):
                pts_val = 0.0

            ms = cell_val(ws.cell(r, cfg["milestone"])) or cell_val(ws.cell(r, cfg["status"])) or "-"

            dac_raw = cell_val(ws.cell(r, cfg.get("dac_month", 9)))
            deliv_raw = cell_val(ws.cell(r, cfg.get("deliv_date", 8)))

            dt_obj = extract_task_date(dac_raw) or extract_task_date(deliv_raw)
            my_label = dt_obj.strftime("%b %Y") if dt_obj else "Unknown Date"

            tasks.append({
                "assignee": v.strip(),
                "task_id": str(tid).strip(),
                "points": pts_val,
                "milestone": str(ms).strip(),
                "month_year": my_label,
                "task_date": dt_obj,
                "source": sname
            })
    return tasks


def filter_tasks_interactively(tasks, pick):
    """Present options: Dump ALL tasks, Single Month & Year, or Date Range / Time Period."""
    print(f"\n  Found {len(tasks)} tasks total for '{pick}'.")
    print("  Select Dump Option:")
    print("    1. Dump ALL tasks (No date filter)")
    print("    2. Dump tasks filtered by a Single Month & Year")
    print("    3. Dump tasks filtered by a Date Range / Time Period (e.g. Jan 2025 to Jun 2026)")

    try:
        choice = input("\n  Enter option (1, 2, or 3, default 1): ").strip()
    except (EOFError, KeyboardInterrupt):
        return tasks

    if choice not in ("2", "3"):
        return tasks

    by_month = {}
    for t in tasks:
        my = t.get("month_year", "Unknown Date")
        by_month.setdefault(my, []).append(t)

    sorted_my = sorted(by_month.keys(), key=parse_month_year_sortkey)

    print(f"\n  Available Month & Year periods for {pick}:")
    for i, my in enumerate(sorted_my, 1):
        cnt = len(by_month[my])
        pts = sum(t["points"] for t in by_month[my])
        print(f"    {i:2d}. {my:<15} ({cnt:2d} tasks | {pts:6.2f} pts)")

    if choice == "2":
        try:
            filter_inp = input("\n  Select Month & Year # or type Month & Year (e.g. Jan 2025): ").strip()
        except (EOFError, KeyboardInterrupt):
            return tasks

        if not filter_inp:
            print("  No filter specified. Using ALL tasks.")
            return tasks

        try:
            idx = int(filter_inp)
            if 1 <= idx <= len(sorted_my):
                selected_my = sorted_my[idx - 1]
                filtered = by_month[selected_my]
                print(f"  Selected: {selected_my} ({len(filtered)} tasks)")
                return filtered
        except ValueError:
            pass

        q = filter_inp.lower()
        matched_tasks = [t for t in tasks if q in t.get("month_year", "").lower()]
        if matched_tasks:
            print(f"  Filtered by '{filter_inp}': {len(matched_tasks)} matching tasks found.")
            return matched_tasks
        else:
            print(f"  No tasks match '{filter_inp}'. Returning all tasks.")
            return tasks

    elif choice == "3":
        print("\n  Enter Time Period Range (e.g. 'Jan 2025 to Jun 2026' or select by numbers):")
        try:
            range_inp = input("  Range or Start Month & Year (e.g. Jan 2025): ").strip()
        except (EOFError, KeyboardInterrupt):
            return tasks

        if not range_inp:
            return tasks

        start_str, end_str = "", ""
        for sep in (" to ", " - ", " : "):
            if sep in range_inp.lower():
                parts = range_inp.lower().split(sep, 1)
                start_str, end_str = parts[0].strip(), parts[1].strip()
                break

        if not start_str:
            start_str = range_inp
            try:
                end_str = input("  End Month & Year   (e.g. Jun 2026): ").strip()
            except (EOFError, KeyboardInterrupt):
                end_str = ""

        start_dt, end_dt = None, None
        try:
            s_idx = int(start_str)
            if 1 <= s_idx <= len(sorted_my):
                start_dt = parse_user_month_year(sorted_my[s_idx - 1])
        except ValueError:
            start_dt = parse_user_month_year(start_str)

        try:
            e_idx = int(end_str)
            if 1 <= e_idx <= len(sorted_my):
                end_dt = parse_user_month_year(sorted_my[e_idx - 1])
        except ValueError:
            end_dt = parse_user_month_year(end_str)

        if not start_dt:
            print(f"  Could not parse start date '{start_str}'. Using ALL tasks.")
            return tasks
        if not end_dt:
            end_dt = start_dt

        if start_dt > end_dt:
            start_dt, end_dt = end_dt, start_dt

        start_cutoff = datetime(start_dt.year, start_dt.month, 1)
        if end_dt.month == 12:
            next_m = datetime(end_dt.year + 1, 1, 1)
        else:
            next_m = datetime(end_dt.year, end_dt.month + 1, 1)
        end_cutoff = next_m - timedelta(seconds=1)

        matched_tasks = []
        for t in tasks:
            t_dt = t.get("task_date")
            if t_dt and start_cutoff <= t_dt <= end_cutoff:
                matched_tasks.append(t)

        start_lbl = start_cutoff.strftime("%b %Y")
        end_lbl = end_cutoff.strftime("%b %Y")
        print(f"\n  Filtered Time Period ({start_lbl} to {end_lbl}): {len(matched_tasks)} matching tasks found.")
        return matched_tasks


def update_master_excel(master_path, tasks, target_name):
    """Create or append task details to master Excel matching exact layout with Period column and deduplication."""
    if not tasks:
        return master_path

    existing_tasks = set()
    if master_path.exists():
        wb = openpyxl.load_workbook(str(master_path))
        ws = wb.active
        current_resource = None
        for r in range(2, ws.max_row + 1):
            cell_name = ws.cell(r, 2).value
            if cell_name and str(cell_name).strip():
                current_resource = str(cell_name).strip().lower()
            cell_task = ws.cell(r, 4).value
            if current_resource and cell_task and str(cell_task).strip():
                existing_tasks.add((current_resource, str(cell_task).strip().lower()))
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mobily billing actuals"

    tasks_to_add = [t for t in tasks if (target_name.lower(), t["task_id"].lower()) not in existing_tasks]

    if len(tasks_to_add) < len(tasks):
        skipped_cnt = len(tasks) - len(tasks_to_add)
        print(f"  Note: Skipped {skipped_cnt} task(s) for '{target_name}' that are already in Master Excel to avoid duplicate rows.")

    if not tasks_to_add:
        print(f"  All selected tasks for '{target_name}' are already present in Master Excel. No new rows added.")
        return master_path

    hdr_fill = PatternFill("solid", fgColor="9BC2E6")
    hdr_font = Font("Calibri", bold=True, size=10, color="000000")
    green_fill = PatternFill("solid", fgColor="00B050")
    green_font = Font("Calibri", bold=True, size=10, color="FFFFFF")
    cell_font = Font("Calibri", size=10, color="000000")
    
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    headers = [
        "SNO", "Resource", "Date onboarded", "Task ID", "Estimator Points",
        "Current Milestone", "Period / Month-Year", "Invoicing calendar", "Invoice",
        "Comments", "Invoice due", "Earned SP\n(basis approved TS till Apr '26)"
    ]

    if ws.cell(1, 1).value is None:
        ws.row_dimensions[1].height = 40
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(1, col_idx, h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = align_center
            cell.border = thin_border

    start_sno = 1
    start_row = 2
    if ws.max_row > 1 and ws.cell(2, 1).value is not None:
        data_rows = ws.max_row
        start_row = data_rows + 1
        sno_count = 0
        for r in range(2, data_rows + 1):
            if ws.cell(r, 1).value is not None:
                sno_count += 1
        start_sno = sno_count + 1

    num_tasks = len(tasks_to_add)
    end_row = start_row + (num_tasks * 3) - 1

    ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
    ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
    ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
    ws.merge_cells(start_row=start_row, start_column=12, end_row=end_row, end_column=12)

    ws.cell(start_row, 1, start_sno).alignment = align_center
    ws.cell(start_row, 2, target_name).alignment = align_left
    ws.cell(start_row, 3, "").alignment = align_center

    earned_sp = ws.cell(start_row, 12, f"=SUM(H{start_row}:H{end_row})")
    earned_sp.alignment = align_right
    earned_sp.number_format = "#,##0.00"

    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = 20
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = cell_font
            cell.border = thin_border

    milestone_stages = ["UAT", "RFS", "PAC"]

    for k, t in enumerate(tasks_to_add):
        t_start = start_row + (k * 3)
        t_end = t_start + 2

        ws.merge_cells(start_row=t_start, start_column=4, end_row=t_end, end_column=4)
        ws.merge_cells(start_row=t_start, start_column=5, end_row=t_end, end_column=5)
        ws.merge_cells(start_row=t_start, start_column=6, end_row=t_end, end_column=6)
        ws.merge_cells(start_row=t_start, start_column=7, end_row=t_end, end_column=7)

        ws.cell(t_start, 4, t["task_id"]).alignment = align_center
        
        pts_cell = ws.cell(t_start, 5, t["points"])
        pts_cell.alignment = align_right
        pts_cell.number_format = "#,##0.00"

        ms_cell = ws.cell(t_start, 6, t["milestone"])
        ms_cell.alignment = align_center
        ms_cell.fill = green_fill
        ms_cell.font = green_font

        my_cell = ws.cell(t_start, 7, t.get("month_year", "-"))
        my_cell.alignment = align_center

        for idx, stage in enumerate(milestone_stages):
            curr_r = t_start + idx
            factor = 0.6 if idx == 0 else 0.2
            
            inv_cell = ws.cell(curr_r, 8, f"={factor}*E{t_start}")
            inv_cell.alignment = align_right
            inv_cell.number_format = "#,##0.00"

            inv_amt = ws.cell(curr_r, 9, f"=H{curr_r}*10451")
            inv_amt.alignment = align_right
            inv_amt.number_format = "#,##0"

            ws.cell(curr_r, 10, stage).alignment = align_left
            ws.cell(curr_r, 11, "").alignment = align_center

    col_widths = {1: 8, 2: 22, 3: 15, 4: 15, 5: 14, 6: 16, 7: 18, 8: 16, 9: 14, 10: 18, 11: 14, 12: 20}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(str(master_path))
    return master_path


def match(query, names):
    """Case-insensitive partial name matching."""
    q = query.lower().strip()
    exact = [n for n in names if n.lower() == q]
    if exact: return exact
    return [n for n in names if q in n.lower()]


def main():
    print("\n  +-----------------------------------------------------+")
    print("  |    TASK & BILLING EXTRACTOR - Master Excel Sync     |")
    print("  +----------------------------------------------------+\n")

    wb, names, src_path = load_data()
    print(f"  {len(names)} assignees, {sum(names.values())} total tasks.\n")
    print(f"  Target Master File: {MASTER_PATH.name}\n")

    while True:
        sorted_n = sorted(names, key=str.lower)
        for i, n in enumerate(sorted_n, 1):
            print(f"  {i:3d}. {n:<40s} ({names[n]:3d} tasks)")

        try:
            inp = input("\n  Name or # (q to quit): ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if inp.lower() in ("q","quit","exit",""):
            break

        pick = None
        try:
            num = int(inp)
            if 1 <= num <= len(sorted_n): pick = sorted_n[num-1]
        except ValueError: pass

        if not pick:
            hits = match(inp, names)
            if not hits:
                print(f"  No match for '{inp}'.\n"); continue
            if len(hits) == 1:
                pick = hits[0]
            else:
                print(f"  Multiple matches:")
                for i, h in enumerate(hits, 1): print(f"    {i}. {h}")
                try:
                    c = int(input("  Pick #: "))
                    pick = hits[c-1] if 1 <= c <= len(hits) else None
                except: pass
                if not pick: print("  Invalid.\n"); continue

        tasks = extract_task_details(wb, pick)
        if not tasks:
            print(f"  No tasks for '{pick}'.\n"); continue

        tasks_to_dump = filter_tasks_interactively(tasks, pick)
        if not tasks_to_dump:
            print(f"  No tasks selected for dumping.\n"); continue

        total_pts = sum(t["points"] for t in tasks_to_dump)

        print(f"\n  Extracted for {pick} ({len(tasks_to_dump)} tasks | {total_pts:.2f} total estimator points):")
        print("  " + "-" * 75)
        print(f"  {'Task ID':<15} | {'Estimator Pts':<13} | {'Milestone Reached':<20} | {'Period / Month-Year':<20}")
        print("  " + "-" * 75)
        for t in tasks_to_dump:
            print(f"  {t['task_id']:<15} | {t['points']:<13.2f} | {t['milestone']:<20} | {t['month_year']:<20}")
        print("  " + "-" * 75)

        path = update_master_excel(MASTER_PATH, tasks_to_dump, pick)
        print(f"\n  Master Excel Updated -> {path.name}")
        print(f"  Path: {path}\n")

        try:
            if input("  Another query? (y/n): ").strip().lower() not in ("y","yes"): break
        except: break
        print()

    print("\n  Done.\n")

if __name__ == "__main__":
    main()

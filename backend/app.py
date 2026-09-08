# app.py - Flask Web Server for Cyberpunk Task Extractor Terminal UI
import os
import sys
from pathlib import Path
from datetime import datetime, timedelta

# Add backend directory to sys.path
BACKEND_DIR = Path(__file__).resolve().parent
BASE_DIR = BACKEND_DIR.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from flask import Flask, render_template, request, jsonify, send_file
from werkzeug.utils import secure_filename
import openpyxl

from task_extractor import (
    load_data, extract_task_details, update_master_excel,
    format_month_year, parse_user_month_year, parse_month_year_sortkey,
    extract_task_date, SOURCE, MASTER_PATH, SHEETS, cell_val
)

FRONTEND_DIR = BASE_DIR / "frontend"

app = Flask(
    __name__,
    static_folder=str(FRONTEND_DIR / "static"),
    template_folder=str(FRONTEND_DIR / "templates")
)

app.config['UPLOAD_FOLDER'] = os.path.join(BASE_DIR, 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50 MB max

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)


def parse_workbook_file(filepath):
    """Load workbook from filepath and return assignee names and task counts."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    names = {}
    for sname, cfg in SHEETS.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        for r in range(cfg["start"], ws.max_row + 1):
            v = cell_val(ws.cell(r, cfg["name"]))
            if v and isinstance(v, str):
                n = v.strip()
                if n and n != "-":
                    names[n] = names.get(n, 0) + 1
    return wb, names


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/load", methods=["GET"])
def api_load():
    """Load default workbook data."""
    try:
        wb, names, target_src = load_data()
        sorted_names = sorted([{"name": n, "count": cnt} for n, cnt in names.items()], key=lambda x: x["name"].lower())
        
        return jsonify({
            "success": True,
            "filename": target_src.name,
            "assignees": sorted_names,
            "total_assignees": len(sorted_names),
            "total_tasks": sum(names.values())
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/upload", methods=["POST"])
def api_upload():
    """Handle custom Excel file upload."""
    try:
        if "file" not in request.files:
            return jsonify({"success": False, "error": "No file part in request"}), 400
        
        file = request.files["file"]
        if file.filename == "":
            return jsonify({"success": False, "error": "No file selected"}), 400
        
        filename = secure_filename(file.filename)
        save_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
        file.save(save_path)
        
        wb, names = parse_workbook_file(save_path)
        sorted_names = sorted([{"name": n, "count": cnt} for n, cnt in names.items()], key=lambda x: x["name"].lower())
        
        return jsonify({
            "success": True,
            "filename": filename,
            "filepath": save_path,
            "assignees": sorted_names,
            "total_assignees": len(sorted_names),
            "total_tasks": sum(names.values())
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/assignee-details", methods=["POST"])
def api_assignee_details():
    """Get all tasks and available Month-Year breakdown for selected assignee."""
    try:
        data = request.get_json() or {}
        target_name = data.get("name")
        filepath = data.get("filepath")
        
        if not target_name:
            return jsonify({"success": False, "error": "Target assignee name required"}), 400
        
        if filepath and os.path.exists(filepath):
            wb = openpyxl.load_workbook(filepath, data_only=True)
        else:
            wb, _, _ = load_data()
            
        tasks = extract_task_details(wb, target_name)
        
        by_month = {}
        for t in tasks:
            my = t.get("month_year", "Unknown Date")
            if my not in by_month:
                by_month[my] = {"count": 0, "points": 0.0, "tasks": []}
            by_month[my]["count"] += 1
            by_month[my]["points"] += t["points"]
            by_month[my]["tasks"].append(t)
            
        sorted_my = sorted(by_month.keys(), key=parse_month_year_sortkey)
        month_summary = []
        for my in sorted_my:
            month_summary.append({
                "month_year": my,
                "count": by_month[my]["count"],
                "points": round(by_month[my]["points"], 2)
            })
            
        serializable_tasks = []
        for t in tasks:
            t_copy = dict(t)
            if "task_date" in t_copy and isinstance(t_copy["task_date"], datetime):
                t_copy["task_date"] = t_copy["task_date"].strftime("%Y-%m-%d")
            serializable_tasks.append(t_copy)
            
        return jsonify({
            "success": True,
            "assignee": target_name,
            "total_tasks": len(tasks),
            "total_points": round(sum(t["points"] for t in tasks), 2),
            "month_periods": month_summary,
            "tasks": serializable_tasks
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/process-dump", methods=["POST"])
def api_process_dump():
    """Filter tasks by selected criteria and dump into Master Excel."""
    try:
        data = request.get_json() or {}
        target_name = data.get("name")
        filepath = data.get("filepath")
        filter_type = data.get("filter_type", "all")
        single_month = data.get("single_month")
        start_month = data.get("start_month")
        end_month = data.get("end_month")
        
        if not target_name:
            return jsonify({"success": False, "error": "Target assignee name required"}), 400

        if filepath and os.path.exists(filepath):
            wb = openpyxl.load_workbook(filepath, data_only=True)
        else:
            wb, _, _ = load_data()

        tasks = extract_task_details(wb, target_name)
        if not tasks:
            return jsonify({"success": False, "error": f"No tasks found for {target_name}"}), 404

        filtered_tasks = tasks

        if filter_type == "single" and single_month:
            q = single_month.strip().lower()
            filtered_tasks = [t for t in tasks if q in t.get("month_year", "").lower()]
            filter_description = f"Single Month: {single_month}"

        elif filter_type == "range" and start_month:
            start_dt = parse_user_month_year(start_month)
            end_dt = parse_user_month_year(end_month) if end_month else start_dt

            if not start_dt:
                return jsonify({"success": False, "error": f"Invalid start month format: '{start_month}'"}), 400
            if not end_dt:
                end_dt = start_dt

            if start_dt > end_dt:
                start_dt, end_dt = end_dt, start_dt

            start_cutoff = datetime(start_dt.year, start_dt.month, 1)
            if end_dt.month == 12:
                next_m = datetime(end_dt.year + 1, 1, 1)
            else:
                next_m = datetime(end_dt.year, end_dt.month + 1, 1)
            end_cutoff = next_m - timedelta(seconds=1)

            filtered_tasks = [t for t in tasks if t.get("task_date") and start_cutoff <= t["task_date"] <= end_cutoff]
            filter_description = f"Date Range: {start_cutoff.strftime('%b %Y')} to {end_cutoff.strftime('%b %Y')}"

        else:
            filter_description = "ALL Tasks (No date filter)"

        if not filtered_tasks:
            return jsonify({
                "success": False,
                "error": f"No matching tasks found for filter criteria: {filter_description}"
            }), 404

        master_output_path = update_master_excel(MASTER_PATH, filtered_tasks, target_name)

        serializable_tasks = []
        for t in filtered_tasks:
            t_copy = dict(t)
            if "task_date" in t_copy and isinstance(t_copy["task_date"], datetime):
                t_copy["task_date"] = t_copy["task_date"].strftime("%Y-%m-%d")
            serializable_tasks.append(t_copy)

        return jsonify({
            "success": True,
            "assignee": target_name,
            "filter_applied": filter_description,
            "task_count": len(filtered_tasks),
            "total_points": round(sum(t["points"] for t in filtered_tasks), 2),
            "master_file": master_output_path.name,
            "tasks": serializable_tasks
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/download-master", methods=["GET"])
def api_download_master():
    """Download the generated Billing_Extracted_Report.xlsx file."""
    if not MASTER_PATH.exists():
        return jsonify({"success": False, "error": "Master Excel file has not been generated yet."}), 404
    
    return send_file(
        str(MASTER_PATH),
        as_attachment=True,
        download_name="Billing_Extracted_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def start_server(port=5000):
    print("\n  +-----------------------------------------------------+")
    print("  |   CYBERPUNK TASK EXTRACTOR - Web Server Starting    |")
    print(f"  |   URL: http://127.0.0.1:{port}                        |")
    print("  +----------------------------------------------------+\n")
    app.run(host="127.0.0.1", port=port, debug=True)


if __name__ == "__main__":
    start_server()

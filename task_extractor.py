# task_extractor.py - Extract Task IDs & Billing Details to Master Excel File
# Usage: python task_extractor.py

# ── CHANGE THIS PATH if you move the Excel file elsewhere ──
FILE_PATH = r"c:\Users\abhin\Downloads\Billing Test\abhinav-input-Jun'26dump.xlsx"
MASTER_FILE_NAME = "Billing_Extracted_Report.xlsx"
# ────────────────────────────────────────────────────────────

import sys
from datetime import datetime
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    try: sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except: pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SOURCE = Path(FILE_PATH)
BASE_DIR = SOURCE.parent
MASTER_PATH = BASE_DIR / MASTER_FILE_NAME

# Sheet layout configuration (exclusively 'Jira Dump')
# Jira Dump: hdr=7, start=8, task=6, name=7, points=12, milestone=10, status=3
SHEETS = {
    "Jira Dump": {"hdr": 7, "start": 8, "task": 6, "name": 7, "points": 12, "milestone": 10, "status": 3},
}


def cell_val(cell):
    """Return cell value, skip formulas."""
    v = cell.value
    return None if isinstance(v, str) and v.startswith("=") else v


def load_data():
    """Load workbook and collect all assignees with their task counts from Jira Dump."""
    if not SOURCE.exists():
        print(f"  ERROR: {SOURCE} not found"); sys.exit(1)

    print(f"  Loading {SOURCE.name} ...")
    wb = openpyxl.load_workbook(str(SOURCE), data_only=True)
    names = {}  # name -> count

    for sname, cfg in SHEETS.items():
        if sname not in wb.sheetnames: continue
        ws = wb[sname]
        for r in range(cfg["start"], ws.max_row + 1):
            v = cell_val(ws.cell(r, cfg["name"]))
            if v and isinstance(v, str):
                n = v.strip()
                if n and n != "-":
                    names[n] = names.get(n, 0) + 1

    return wb, names


def extract_task_details(wb, target):
    """Extract Task IDs, Estimator Points, Milestone & Status for target name from Jira Dump."""
    tasks = []
    for sname, cfg in SHEETS.items():
        if sname not in wb.sheetnames: continue
        ws = wb[sname]
        for r in range(cfg["start"], ws.max_row + 1):
            v = cell_val(ws.cell(r, cfg["name"]))
            if not (v and isinstance(v, str) and v.strip().lower() == target.lower()):
                continue
            
            tid = cell_val(ws.cell(r, cfg["task"]))
            if tid is None: continue
            
            pts = cell_val(ws.cell(r, cfg["points"]))
            try:
                pts_val = float(pts) if pts is not None else 0.0
            except (ValueError, TypeError):
                pts_val = 0.0

            ms = cell_val(ws.cell(r, cfg["milestone"])) or cell_val(ws.cell(r, cfg["status"])) or "-"

            tasks.append({
                "assignee": v.strip(),
                "task_id": str(tid).strip(),
                "points": pts_val,
                "milestone": str(ms).strip(),
                "source": sname
            })
    return tasks


def update_master_excel(master_path, tasks, target_name):
    """Create or append task details to master Excel matching exact layout from template image."""
    if not tasks:
        return master_path

    if master_path.exists():
        wb = openpyxl.load_workbook(str(master_path))
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mobily billing actuals"

    # Header styling (light blue fill, dark bold text, wrap text)
    hdr_fill = PatternFill("solid", fgColor="9BC2E6")
    hdr_font = Font("Calibri", bold=True, size=10, color="000000")
    
    # Green milestone styling (matching template image)
    green_fill = PatternFill("solid", fgColor="00B050")
    green_font = Font("Calibri", bold=True, size=10, color="FFFFFF")

    cell_font = Font("Calibri", size=10, color="000000")
    
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    headers = [
        "SNO", "Resource", "Date onboarded", "Task ID", "Estimator Points",
        "Current Milestone", "Invoicing calendar", "Invoice", "Comments",
        "Invoice due", "Earned SP\n(basis approved TS till Apr '26)"
    ]

    # Write headers if sheet is empty
    if ws.cell(1, 1).value is None:
        ws.row_dimensions[1].height = 40
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(1, col_idx, h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = align_center
            cell.border = thin_border

    # Determine starting SNO and next row index
    start_sno = 1
    start_row = 2
    if ws.max_row > 1 and ws.cell(2, 1).value is not None:
        data_rows = ws.max_row
        start_row = data_rows + 1
        sno_count = 0
        for r in range(2, data_rows + 1):
            if ws.cell(r, 1).value is not None:
                sno_count += 1
        start_sno = sno_count + 1

    num_tasks = len(tasks)
    end_row = start_row + (num_tasks * 3) - 1

    # Merge SNO, Resource Name, Date onboarded, and Earned SP vertically for this resource
    ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
    ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
    ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
    ws.merge_cells(start_row=start_row, start_column=11, end_row=end_row, end_column=11)

    ws.cell(start_row, 1, start_sno).alignment = align_center
    ws.cell(start_row, 2, target_name).alignment = align_left
    ws.cell(start_row, 3, "").alignment = align_center

    # Earned SP formula
    earned_sp = ws.cell(start_row, 11, f"=SUM(G{start_row}:G{end_row})")
    earned_sp.alignment = align_right
    earned_sp.number_format = "#,##0.00"

    # Apply font and borders to all cells in the block
    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = 20
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = cell_font
            cell.border = thin_border

    # Format each task (3 sub-rows per task)
    milestone_stages = ["UAT", "RFS", "PAC"]

    for k, t in enumerate(tasks):
        t_start = start_row + (k * 3)
        t_end = t_start + 2

        # Merge Task ID, Estimator Points, Current Milestone across 3 sub-rows
        ws.merge_cells(start_row=t_start, start_column=4, end_row=t_end, end_column=4)
        ws.merge_cells(start_row=t_start, start_column=5, end_row=t_end, end_column=5)
        ws.merge_cells(start_row=t_start, start_column=6, end_row=t_end, end_column=6)

        ws.cell(t_start, 4, t["task_id"]).alignment = align_center
        
        pts_cell = ws.cell(t_start, 5, t["points"])
        pts_cell.alignment = align_right
        pts_cell.number_format = "#,##0.00"

        ms_cell = ws.cell(t_start, 6, t["milestone"])
        ms_cell.alignment = align_center
        ms_cell.fill = green_fill
        ms_cell.font = green_font

        # 3 sub-rows for invoicing calendar, invoice amount, and comments
        for idx, stage in enumerate(milestone_stages):
            curr_r = t_start + idx
            factor = 0.6 if idx == 0 else 0.2
            
            # Invoicing calendar formula e.g. =0.6*E2
            inv_cell = ws.cell(curr_r, 7, f"={factor}*E{t_start}")
            inv_cell.alignment = align_right
            inv_cell.number_format = "#,##0.00"

            # Invoice amount formula e.g. =G2*10451
            inv_amt = ws.cell(curr_r, 8, f"=G{curr_r}*10451")
            inv_amt.alignment = align_right
            inv_amt.number_format = "#,##0"

            # Comments
            ws.cell(curr_r, 9, stage).alignment = align_left
            
            # Invoice due
            ws.cell(curr_r, 10, "").alignment = align_center

    # Column widths
    col_widths = {1: 8, 2: 22, 3: 15, 4: 15, 5: 14, 6: 16, 7: 16, 8: 14, 9: 18, 10: 14, 11: 20}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(str(master_path))
    return master_path


def match(query, names):
    """Case-insensitive partial name matching."""
    q = query.lower().strip()
    exact = [n for n in names if n.lower() == q]
    if exact: return exact
    return [n for n in names if q in n.lower()]


def main():
    print("\n  +-----------------------------------------------------+")
    print("  |    TASK & BILLING EXTRACTOR - Master Excel Sync     |")
    print("  +----------------------------------------------------+\n")

    wb, names = load_data()
    print(f"  {len(names)} assignees, {sum(names.values())} total tasks.\n")
    print(f"  Target Master File: {MASTER_PATH.name}\n")

    while True:
        sorted_n = sorted(names, key=str.lower)
        for i, n in enumerate(sorted_n, 1):
            print(f"  {i:3d}. {n:<40s} ({names[n]:3d} tasks)")

        try:
            inp = input("\n  Name or # (q to quit): ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if inp.lower() in ("q","quit","exit",""):
            break

        pick = None
        try:
            num = int(inp)
            if 1 <= num <= len(sorted_n): pick = sorted_n[num-1]
        except ValueError: pass

        if not pick:
            hits = match(inp, names)
            if not hits:
                print(f"  No match for '{inp}'.\n"); continue
            if len(hits) == 1:
                pick = hits[0]
            else:
                print(f"  Multiple matches:")
                for i, h in enumerate(hits, 1): print(f"    {i}. {h}")
                try:
                    c = int(input("  Pick #: "))
                    pick = hits[c-1] if 1 <= c <= len(hits) else None
                except: pass
                if not pick: print("  Invalid.\n"); continue

        tasks = extract_task_details(wb, pick)
        if not tasks:
            print(f"  No tasks for '{pick}'.\n"); continue

        total_pts = sum(t["points"] for t in tasks)

        # Print Task IDs & Points
        print(f"\n  Extracted for {pick} ({len(tasks)} tasks | {total_pts:.2f} total estimator points):")
        print("  " + "-" * 60)
        print(f"  {'Task ID':<15} | {'Estimator Pts':<13} | {'Milestone Reached':<20}")
        print("  " + "-" * 60)
        for t in tasks:
            print(f"  {t['task_id']:<15} | {t['points']:<13.2f} | {t['milestone']:<20}")
        print("  " + "-" * 60)

        # Save to Master Excel
        path = update_master_excel(MASTER_PATH, tasks, pick)
        print(f"\n  Master Excel Updated -> {path.name}")
        print(f"  Path: {path}\n")

        try:
            if input("  Another query? (y/n): ").strip().lower() not in ("y","yes"): break
        except: break
        print()

    print("\n  Done.\n")

if __name__ == "__main__":
    main()



